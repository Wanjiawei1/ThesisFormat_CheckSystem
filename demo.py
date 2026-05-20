from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from datetime import datetime
import zipfile
import tempfile
import shutil
import re


_comment_counter = 0
_comments = {}  # {comment_id: comment_text}
_author = "Reviewer"  # 全局作者名，由 add_comment / add_comment_to_text 设置

def add_comment(paragraph, comment_text, author="Reviewer"):
    """对整个段落添加批注。

    Args:
        paragraph: docx.paragraphs[i]，目标段落对象
        comment_text: 批注内容
        author: 批注作者名，默认 "Reviewer"

    Returns:
        comment_id (str)

    Example:
        add_comment(doc.paragraphs[0], "这段需要修改")
    """
    global _comment_counter, _author
    _author = author
    comment_id = str(_comment_counter)
    _comment_counter += 1

    # commentRangeStart 插入到 pPr 之后
    insert_idx = 0
    for i, child in enumerate(paragraph._p):
        if child.tag.endswith('pPr'):
            insert_idx = i + 1
            break
    range_start = OxmlElement('w:commentRangeStart')
    range_start.set(qn('w:id'), comment_id)
    paragraph._p.insert(insert_idx, range_start)

    # commentRangeEnd 插入到段落末尾
    range_end = OxmlElement('w:commentRangeEnd')
    range_end.set(qn('w:id'), comment_id)
    paragraph._p.append(range_end)

    # commentReference 放在单独的 run 中
    ref_run = paragraph.add_run()
    r_ref = OxmlElement('w:commentReference')
    r_ref.set(qn('w:id'), comment_id)
    ref_run._r.append(r_ref)

    _comments[comment_id] = comment_text
    return comment_id


def _clone_rpr(run_element):
    """复制 run 的格式属性"""
    for child in run_element:
        if child.tag.endswith('rPr'):
            from copy import deepcopy
            return deepcopy(child)
    return None


def _make_run(text, rpr=None):
    """创建一个带可选格式的 run 元素"""
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def add_comment_to_text(paragraph, search_text, comment_text, author="Reviewer"):
    """对段落中的特定文字添加批注（支持跨 run 搜索）。

    Word 经常将一句话切成多个 run，本函数先拼接全文定位索引，
    再映射回各 run 进行精确拆分，确保只批注目标文字。

    Args:
        paragraph: docx.paragraphs[i]，目标段落对象
        search_text: 要批注的文字（可跨越多个 run）
        comment_text: 批注内容
        author: 批注作者名，默认 "Reviewer"

    Returns:
        comment_id (str)。如果 search_text 未找到，打印警告并退回到整段批注。

    Example:
        add_comment_to_text(doc.paragraphs[3], "浙江工业大学", "请确认学校名称")
    """
    global _comment_counter, _author

    # 收集所有 run 的信息：(run元素, 文本内容, 在段落子元素列表中的索引)
    run_infos = []
    for idx, child in enumerate(paragraph._p):
        if not child.tag.endswith('r'):
            continue
        t_el = child.find(qn('w:t'))
        text = t_el.text if t_el is not None and t_el.text else ''
        run_infos.append((child, text, idx))

    full_text = ''.join(text for _, text, _ in run_infos)

    if search_text not in full_text:
        print(f"警告: 未找到文字 \"{search_text}\"，退回到整段批注")
        return add_comment(paragraph, comment_text, author)

    _author = author
    comment_id = str(_comment_counter)
    _comment_counter += 1

    start_pos = full_text.index(search_text)
    end_pos = start_pos + len(search_text)

    # 计算每个 run 的字符偏移范围，确定需要拆分的 run
    replacements = []  # [(段落子元素索引, 要插入的新元素列表)]
    current_offset = 0

    for run_el, text, p_idx in run_infos:
        run_start = current_offset
        run_end = current_offset + len(text)
        rpr = _clone_rpr(run_el)

        # 这个 run 与目标范围没有交集，跳过
        if run_end <= start_pos or run_start >= end_pos:
            current_offset += len(text)
            continue

        parts = []

        # run 前缀（在目标范围之前的部分）
        if run_start < start_pos:
            prefix = text[:start_pos - run_start]
            parts.append(_make_run(prefix, _clone_rpr(run_el) if rpr else None))

        # run 中属于目标范围的开始处插入 commentRangeStart
        if run_start <= start_pos < run_end:
            parts.append(OxmlElement('w:commentRangeStart'))
            parts[-1].set(qn('w:id'), comment_id)

        # run 中属于目标范围的文本
        seg_start = max(0, start_pos - run_start)
        seg_end = min(len(text), end_pos - run_start)
        parts.append(_make_run(text[seg_start:seg_end], _clone_rpr(run_el) if rpr else None))

        # run 中目标范围的结束处插入 commentRangeEnd + commentReference
        if run_start < end_pos <= run_end:
            parts.append(OxmlElement('w:commentRangeEnd'))
            parts[-1].set(qn('w:id'), comment_id)
            ref_run_el = OxmlElement('w:r')
            if rpr:
                ref_run_el.append(rpr)
            r_ref = OxmlElement('w:commentReference')
            r_ref.set(qn('w:id'), comment_id)
            ref_run_el.append(r_ref)
            parts.append(ref_run_el)

        # run 后缀（在目标范围之后的部分）
        if run_end > end_pos:
            suffix = text[end_pos - run_start:]
            parts.append(_make_run(suffix, _clone_rpr(run_el) if rpr else None))

        replacements.append((p_idx, parts))
        current_offset += len(text)

    # 按索引从大到小排序，避免插入位置偏移
    replacements.sort(key=lambda x: x[0], reverse=True)

    for p_idx, parts in replacements:
        old_el = list(paragraph._p)[p_idx]
        paragraph._p.remove(old_el)
        for offset, el in enumerate(parts):
            paragraph._p.insert(p_idx + offset, el)

    _comments[comment_id] = comment_text
    return comment_id


def save_with_comments(doc, output_path):
    """保存文档并将所有已添加的批注写入文件。

    必须在 add_comment / add_comment_to_text 之后调用。

    Args:
        doc: Document 对象
        output_path: 输出文件路径

    Example:
        save_with_comments(doc, "output.docx")
    """
    global _comment_counter, _comments, _author
    author = _author
    W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    doc.save(output_path)

    comments = _comments

    # 构建 comments.xml
    xml_parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:comments xmlns:w="' + W_NS + '">',
    ]
    for cid, text in comments.items():
        date_str = datetime.now().strftime('%Y-%m-%dT%H:%M:%SZ')
        xml_parts.append(
            f'  <w:comment w:id="{cid}" w:author="{author}" '
            f'w:initials="{author[0]}" w:date="{date_str}">'
        )
        xml_parts.append('    <w:p><w:r>')
        xml_parts.append(f'      <w:t>{text}</w:t>')
        xml_parts.append('    </w:r></w:p>')
        xml_parts.append('  </w:comment>')
    xml_parts.append('</w:comments>')
    comments_xml = '\n'.join(xml_parts).encode('utf-8')

    # ZIP 后处理：替换 comments.xml，更新 Content_Types 和属性
    with zipfile.ZipFile(output_path, 'r') as zin:
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp_path = tmp.name
        with zipfile.ZipFile(tmp_path, 'w') as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                # 跳过旧的 comments.xml
                if item.filename == 'word/comments.xml':
                    continue
                # 更新 Content_Types.xml
                elif item.filename == '[Content_Types].xml':
                    from lxml import etree
                    types_root = etree.fromstring(data)
                    ns = {'ct': 'http://schemas.openxmlformats.org/package/2006/content-types'}
                    if not types_root.xpath(
                        "//ct:Override[@PartName='/word/comments.xml']",
                        namespaces=ns
                    ):
                        override = etree.Element('Override')
                        override.set('PartName', '/word/comments.xml')
                        override.set('ContentType',
                            'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml')
                        types_root.append(override)
                    zout.writestr(item,
                        etree.tostring(types_root, encoding='utf-8', xml_declaration=True))
                # 更新作者名
                elif item.filename == 'docProps/core.xml':
                    from lxml import etree
                    core_root = etree.fromstring(data)
                    ns = {
                        'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
                        'dc': 'http://purl.org/dc/elements/1.1/'
                    }
                    for el in core_root.xpath('//dc:creator', namespaces=ns):
                        el.text = author
                    for el in core_root.xpath('//cp:lastModifiedBy', namespaces=ns):
                        el.text = author
                    zout.writestr(item,
                        etree.tostring(core_root, encoding='utf-8', xml_declaration=True))
                # 确保 document.xml.rels 包含 comments 关系
                elif item.filename == 'word/_rels/document.xml.rels':
                    from lxml import etree
                    rels_root = etree.fromstring(data)
                    ns = {'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}
                    if not rels_root.xpath(
                        "//r:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments']",
                        namespaces=ns
                    ):
                        # 找到最大 rId
                        max_id = 0
                        for rid in rels_root.xpath("//r:Relationship/@Id", namespaces=ns):
                            m = re.search(r'rId(\d+)', rid)
                            if m:
                                max_id = max(max_id, int(m.group(1)))
                        new_rel = etree.Element('Relationship')
                        new_rel.set('Id', f'rId{max_id + 1}')
                        new_rel.set('Type',
                            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments')
                        new_rel.set('Target', 'comments.xml')
                        rels_root.append(new_rel)
                    zout.writestr(item,
                        etree.tostring(rels_root, encoding='utf-8', xml_declaration=True))
                else:
                    zout.writestr(item, data)
            # 写入新的 comments.xml
            zout.writestr('word/comments.xml', comments_xml)
    shutil.move(tmp_path, output_path)


# --- 测试运行 ---
file_path = "paper.docx"
doc = Document(file_path)

# 1. 整段批注
add_comment(doc.paragraphs[8], "开头可以使用更正式的称谓")

# 2. 特定文字批注
add_comment_to_text(doc.paragraphs[1], "个性化学习路径", "这个术语需要统一定义")

save_with_comments(doc, "paper_with_comments.docx")
print("保存成功！请检查 paper_with_comments.docx")
"""
高亮显示文档中的格式错误
功能：
1. 复制原始Word文档
2. 根据检测结果在错误位置添加高亮背景
3. 生成带有高亮的新文档
"""

import shutil
from docx import Document
import os


class HighlightAdder:
    """高亮添加器"""
    
    def __init__(self):
        pass
    
    def copy_file(self, original_file, target_file):
        """复制文件"""
        try:
            shutil.copyfile(original_file, target_file)
            return True
        except Exception as e:
            print(f"复制文件失败: {str(e)}")
            return False
    
    def highlight_paragraph(self, paragraph):
        """在段落中添加高亮背景"""
        # 为段落中的每个run添加高亮背景
        for run in paragraph.runs:
            # 设置背景色为黄色
            run.font.highlight_color = 7  # 7 表示黄色
    
    def add_highlights_from_results(self, original_file, output_file, issues):
        """根据检测结果添加高亮"""
        # 复制文件
        if not self.copy_file(original_file, output_file):
            return False
        
        try:
            # 打开副本
            doc = Document(output_file)
            
            # 遍历所有问题
            for issue in issues:
                paragraph_index = issue.get('paragraph_index', 0)
                if 0 < paragraph_index <= len(doc.paragraphs):
                    # 获取目标段落
                    target_paragraph = doc.paragraphs[paragraph_index - 1]
                    # 添加高亮背景
                    self.highlight_paragraph(target_paragraph)
            
            # 保存修改
            doc.save(output_file)
            return True
        except Exception as e:
            print(f"添加高亮失败: {str(e)}")
            return False


def main():
    """主函数"""
    import argparse
    
    # 解析命令行参数
    parser = argparse.ArgumentParser(description='为Word文档添加格式问题高亮')
    parser.add_argument('input_file', help='输入Word文档路径')
    parser.add_argument('-o', '--output', help='输出文档路径（默认：原文件名+_高亮版）', default=None)
    args = parser.parse_args()
    
    # 检查输入文件是否存在
    if not os.path.exists(args.input_file):
        print(f"错误：文件不存在: {args.input_file}")
        return
    
    # 生成输出路径
    if not args.output:
        base_name = os.path.basename(args.input_file)
        name_without_ext = os.path.splitext(base_name)[0]
        # 将输出文件保存到当前目录
        output_path = os.path.join(os.getcwd(), f"{name_without_ext}_高亮版.docx")
    else:
        output_path = args.output
    
    print(f"正在处理文件: {args.input_file}")
    print(f"输出文件: {output_path}")
    
    # 创建高亮添加器
    highlight_adder = HighlightAdder()
    
    # 这里需要从其他检测模块获取结果
    # 暂时使用空列表作为示例
    issues = []
    
    # 添加高亮
    if highlight_adder.add_highlights_from_results(args.input_file, output_path, issues):
        print(f"\n[通过] 带高亮的文档已生成: {output_path}")
    else:
        print("[错误] 添加高亮失败")


if __name__ == "__main__":
    main()
"""
Word 文档批注工具模块

功能：
1. 底层批注操作（add_comment / add_comment_to_text / save_with_comments）
2. 各 checker 的错误收集器（collect_xxx_errors）
3. 统一入口 run_all_checks_with_comments()
"""

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from datetime import datetime
import zipfile
import tempfile
import shutil
import re
import copy
import os


# ============================================================
# 底层批注操作
# ============================================================

_comment_counter = 0
_comments = {}  # {comment_id: comment_text}
_author = "格式检测器"


def reset_comments(author="格式检测器"):
    """重置批注状态（每次新文档前调用）"""
    global _comment_counter, _comments, _author
    _comment_counter = 0
    _comments = {}
    _author = author


def add_comment(paragraph, comment_text, author=None):
    """对整个段落添加批注。

    Args:
        paragraph: docx.paragraphs[i]，目标段落对象
        comment_text: 批注内容
        author: 批注作者名，默认使用全局 _author

    Returns:
        comment_id (str)
    """
    global _comment_counter, _author
    if author:
        _author = author
    comment_id = str(_comment_counter)
    _comment_counter += 1

    # commentRangeStart 插入到 pPr 之后
    insert_idx = 0
    for i, child in enumerate(paragraph._p):
        if child.tag.endswith('pPr'):
            insert_idx = i + 1
            break
    range_start = OxmlElement('w:commentRangeStart')
    range_start.set(qn('w:id'), comment_id)
    paragraph._p.insert(insert_idx, range_start)

    # commentRangeEnd 插入到段落末尾
    range_end = OxmlElement('w:commentRangeEnd')
    range_end.set(qn('w:id'), comment_id)
    paragraph._p.append(range_end)

    # commentReference 放在单独的 run 中
    ref_run = paragraph.add_run()
    r_ref = OxmlElement('w:commentReference')
    r_ref.set(qn('w:id'), comment_id)
    ref_run._r.append(r_ref)

    _comments[comment_id] = comment_text
    return comment_id


def _clone_rpr(run_element):
    """复制 run 的格式属性"""
    for child in run_element:
        if child.tag.endswith('rPr'):
            return copy.deepcopy(child)
    return None


def _make_run(text, rpr=None):
    """创建一个带可选格式的 run 元素"""
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def add_comment_to_text(paragraph, search_text, comment_text, author=None):
    """对段落中的特定文字添加批注（支持跨 run 搜索）。

    Args:
        paragraph: docx.paragraphs[i]，目标段落对象
        search_text: 要批注的文字（可跨越多个 run）
        comment_text: 批注内容
        author: 批注作者名，默认使用全局 _author

    Returns:
        comment_id (str)。如果 search_text 未找到，退回到整段批注。
    """
    global _comment_counter, _author
    if author:
        _author = author

    # 收集所有 run 的信息
    run_infos = []
    for idx, child in enumerate(paragraph._p):
        if not child.tag.endswith('r'):
            continue
        t_el = child.find(qn('w:t'))
        text = t_el.text if t_el is not None and t_el.text else ''
        run_infos.append((child, text, idx))

    full_text = ''.join(text for _, text, _ in run_infos)

    if search_text not in full_text:
        print(f"警告: 未找到文字 \"{search_text}\"，退回到整段批注")
        return add_comment(paragraph, comment_text, author)

    comment_id = str(_comment_counter)
    _comment_counter += 1

    start_pos = full_text.index(search_text)
    end_pos = start_pos + len(search_text)

    replacements = []
    current_offset = 0

    for run_el, text, p_idx in run_infos:
        run_start = current_offset
        run_end = current_offset + len(text)
        rpr = _clone_rpr(run_el)

        if run_end <= start_pos or run_start >= end_pos:
            current_offset += len(text)
            continue

        parts = []

        if run_start < start_pos:
            prefix = text[:start_pos - run_start]
            parts.append(_make_run(prefix, _clone_rpr(run_el) if rpr else None))

        if run_start <= start_pos < run_end:
            parts.append(OxmlElement('w:commentRangeStart'))
            parts[-1].set(qn('w:id'), comment_id)

        seg_start = max(0, start_pos - run_start)
        seg_end = min(len(text), end_pos - run_start)
        parts.append(_make_run(text[seg_start:seg_end], _clone_rpr(run_el) if rpr else None))

        if run_start < end_pos <= run_end:
            parts.append(OxmlElement('w:commentRangeEnd'))
            parts[-1].set(qn('w:id'), comment_id)
            ref_run_el = OxmlElement('w:r')
            if rpr:
                ref_run_el.append(rpr)
            r_ref = OxmlElement('w:commentReference')
            r_ref.set(qn('w:id'), comment_id)
            ref_run_el.append(r_ref)
            parts.append(ref_run_el)

        if run_end > end_pos:
            suffix = text[end_pos - run_start:]
            parts.append(_make_run(suffix, _clone_rpr(run_el) if rpr else None))

        replacements.append((p_idx, parts))
        current_offset += len(text)

    replacements.sort(key=lambda x: x[0], reverse=True)

    for p_idx, parts in replacements:
        old_el = list(paragraph._p)[p_idx]
        paragraph._p.remove(old_el)
        for offset, el in enumerate(parts):
            paragraph._p.insert(p_idx + offset, el)

    _comments[comment_id] = comment_text
    return comment_id


def save_with_comments(doc, output_path):
    """保存文档并将所有已添加的批注写入文件。

    必须在 add_comment / add_comment_to_text 之后调用。

    Args:
        doc: Document 对象
        output_path: 输出文件路径
    """
    global _comment_counter, _comments, _author
    author = _author
    W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    doc.save(output_path)

    comments = _comments

    # 构建 comments.xml
    xml_parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:comments xmlns:w="' + W_NS + '">',
    ]
    for cid, text in comments.items():
        date_str = datetime.now().strftime('%Y-%m-%dT%H:%M:%SZ')
        xml_parts.append(
            f'  <w:comment w:id="{cid}" w:author="{author}" '
            f'w:initials="{author[0]}" w:date="{date_str}">'
        )
        xml_parts.append('    <w:p><w:r>')
        xml_parts.append(f'      <w:t>{text}</w:t>')
        xml_parts.append('    </w:r></w:p>')
        xml_parts.append('  </w:comment>')
    xml_parts.append('</w:comments>')
    comments_xml = '\n'.join(xml_parts).encode('utf-8')

    # ZIP 后处理
    with zipfile.ZipFile(output_path, 'r') as zin:
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp_path = tmp.name
        with zipfile.ZipFile(tmp_path, 'w') as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == 'word/comments.xml':
                    continue
                elif item.filename == '[Content_Types].xml':
                    from lxml import etree
                    types_root = etree.fromstring(data)
                    ns = {'ct': 'http://schemas.openxmlformats.org/package/2006/content-types'}
                    if not types_root.xpath(
                        "//ct:Override[@PartName='/word/comments.xml']",
                        namespaces=ns
                    ):
                        override = etree.Element('Override')
                        override.set('PartName', '/word/comments.xml')
                        override.set('ContentType',
                            'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml')
                        types_root.append(override)
                    zout.writestr(item,
                        etree.tostring(types_root, encoding='utf-8', xml_declaration=True))
                elif item.filename == 'docProps/core.xml':
                    from lxml import etree
                    core_root = etree.fromstring(data)
                    ns = {
                        'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
                        'dc': 'http://purl.org/dc/elements/1.1/'
                    }
                    for el in core_root.xpath('//dc:creator', namespaces=ns):
                        el.text = author
                    for el in core_root.xpath('//cp:lastModifiedBy', namespaces=ns):
                        el.text = author
                    zout.writestr(item,
                        etree.tostring(core_root, encoding='utf-8', xml_declaration=True))
                elif item.filename == 'word/_rels/document.xml.rels':
                    from lxml import etree
                    rels_root = etree.fromstring(data)
                    ns = {'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}
                    if not rels_root.xpath(
                        "//r:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments']",
                        namespaces=ns
                    ):
                        max_id = 0
                        for rid in rels_root.xpath("//r:Relationship/@Id", namespaces=ns):
                            m = re.search(r'rId(\d+)', rid)
                            if m:
                                max_id = max(max_id, int(m.group(1)))
                        new_rel = etree.Element('Relationship')
                        new_rel.set('Id', f'rId{max_id + 1}')
                        new_rel.set('Type',
                            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments')
                        new_rel.set('Target', 'comments.xml')
                        rels_root.append(new_rel)
                    zout.writestr(item,
                        etree.tostring(rels_root, encoding='utf-8', xml_declaration=True))
                else:
                    zout.writestr(item, data)
            # 写入新的 comments.xml
            zout.writestr('word/comments.xml', comments_xml)
    shutil.move(tmp_path, output_path)


def collect_and_add_comments(doc, issues):
    """统一的批注添加函数。

    Args:
        doc: Document 对象
        issues: list of dict, 每个 dict 包含:
            - paragraph_index: int (0-based 段落索引)
            - comment: str (批注内容)
            - search_text: str (可选，对特定文字添加批注)

    Returns:
        int: 成功添加的批注数量
    """
    count = 0
    paragraphs = doc.paragraphs
    for issue in issues:
        idx = issue.get('paragraph_index')
        comment = issue.get('comment', '')
        search = issue.get('search_text')

        if idx is None or idx < 0 or idx >= len(paragraphs):
            continue
        if not comment:
            continue

        para = paragraphs[idx]
        try:
            if search:
                add_comment_to_text(para, search, comment)
            else:
                add_comment(para, comment)
            count += 1
        except Exception as e:
            print(f"  [警告] 添加批注失败 (段落{idx}): {e}")
    return count


# ============================================================
# 各 checker 的错误收集器
# 返回 list of dict: {paragraph_index, comment, search_text?}
# ============================================================

def _find_paragraph_by_text(doc, search_text):
    """在文档中查找包含指定文本的段落索引"""
    for i, para in enumerate(doc.paragraphs):
        if search_text in para.text:
            return i
    return None


def _find_cover_field_paragraph(doc, field, cover_paras):
    """在封面区域查找特定字段的段落索引"""
    keyword_map = {
        '学校名称': ['浙江工业大学'],
        '论文题目': ['题目'],
        '学院': ['学院'],
        '班级': ['班级'],
        '学号': ['学号'],
        '学生姓名': ['姓名', '学生姓名'],
        '指导老师': ['指导老师', '指导教师'],
        '专业': ['专业'],
        '提交日期': ['提交日期', '年', '月'],
    }
    keywords = keyword_map.get(field, [])
    for idx in cover_paras:
        text = doc.paragraphs[idx].text.strip()
        for kw in keywords:
            if kw in text:
                return idx
    return cover_paras[0] if cover_paras else 0


def _find_ref_paragraph(refs, ref_num):
    """根据文献编号查找段落索引"""
    for ref in refs:
        if ref['number'] == ref_num:
            return ref['paragraph_index']
    return None


def collect_ack_appendix_errors(docx_path):
    """收集致谢 & 附录检测错误（附录不存在不算错误）"""
    from .ack_appendix_checker import AckAppendixChecker
    checker = AckAppendixChecker()
    results = checker.check_ack_and_appendix(docx_path)
    issues = []

    ack = results.get('acknowledgment')
    if ack:
        for p in ack.get('title_problems', []):
            issues.append({'paragraph_index': ack['title_index'], 'comment': f"【致谢标题】{p}"})
        for bi in ack.get('body_issues', []):
            for p in bi['problems']:
                issues.append({'paragraph_index': bi['paragraph_index'], 'comment': f"【致谢正文】{p}"})

    # 附录不存在不算错误，只记录有附录时的格式问题
    for appendix in results.get('appendices', []):
        for p in appendix.get('title_problems', []):
            issues.append({'paragraph_index': appendix['title_index'], 'comment': f"【附录标题】{p}"})
        for bi in appendix.get('body_issues', []):
            for p in bi['problems']:
                issues.append({'paragraph_index': bi['paragraph_index'], 'comment': f"【附录正文】{p}"})

    return issues


def collect_body_text_errors(docx_path):
    """收集正文格式检测错误"""
    from .body_text_checker import BodyTextChecker
    checker = BodyTextChecker()
    results = checker.check_body_text(docx_path)
    issues = []

    for sample in results.get('title_samples', []):
        for issue_text in sample['issues']:
            issues.append({'paragraph_index': sample['paragraph_index'], 'comment': f"【标题格式】{issue_text}"})

    for ai in results.get('all_body_issues', []):
        para_idx = ai['paragraph_index'] - 1  # 1-based -> 0-based
        for issue_text in ai['issues']:
            comment = f"【正文格式】{issue_text}"
            if ai.get('chapter_position'):
                comment += f"\n位置: {ai['chapter_position']}"
            issues.append({'paragraph_index': para_idx, 'comment': comment})

    return issues


def collect_chapter_title_errors(docx_path):
    """收集章节标题检测错误"""
    from .chapter_title_checker import ChapterTitleChecker
    checker = ChapterTitleChecker()
    titles, format_issues, continuity_issues = checker.analyze_titles(docx_path)
    issues = []

    for issue in format_issues:
        issues.append({
            'paragraph_index': issue['paragraph_index'],
            'comment': f"【章节标题】{issue['description']}"
        })

    for issue in continuity_issues:
        issues.append({
            'paragraph_index': issue['paragraph_index'],
            'comment': f"【标题编号】{issue['description']}"
        })

    return issues


def collect_commitment_errors(docx_path):
    """收集承诺书检测错误"""
    from .commitment import check_commitment_info_by_zip
    has_commitment, has_jpeg, has_date = check_commitment_info_by_zip(docx_path)
    issues = []

    if not has_commitment:
        issues.append({'paragraph_index': 0, 'comment': "【承诺书】未检测到承诺书内容（关键词：诚信承诺书/本人慎重承诺/签名）"})
    else:
        if not has_jpeg:
            issues.append({'paragraph_index': 0, 'comment': "【承诺书】未检测到 JPEG 图片（可能是签名图片）"})
        if not has_date:
            issues.append({'paragraph_index': 0, 'comment': "【承诺书】未检测到日期填写（格式如：2024年X月X日）"})

    return issues


def collect_cover_page_errors(docx_path):
    """收集封面格式检测错误"""
    from .CoverPage import check_cover_page_final
    result = check_cover_page_final(docx_path)
    doc = Document(docx_path)
    issues = []

    cover_paras = [i for i, p in enumerate(doc.paragraphs) if p.text.strip()][:30]

    for field, status in result.items():
        if not status.get('内容') and field != '校外指导教师':
            target_idx = cover_paras[0] if cover_paras else 0
            issues.append({'paragraph_index': target_idx, 'comment': f"【封面】{field}：内容未填写或未检测到"})
        if status.get('样式') is False:
            target_idx = _find_cover_field_paragraph(doc, field, cover_paras)
            reason = status.get('原因', '格式不符合要求')
            issues.append({'paragraph_index': target_idx, 'comment': f"【封面】{field}：{reason}"})

    return issues


def collect_figure_table_errors(docx_path):
    """收集图表公式术语检测错误"""
    from .figure_table_formula_term_checker import FigureTableFormulaTermChecker
    checker = FigureTableFormulaTermChecker()
    results = checker.check_all(docx_path)
    doc = Document(docx_path)
    issues = []

    for fig in results.get('figures', []):
        for issue_text in fig['info']['issues']:
            para_idx = _find_paragraph_by_text(doc, f"图{fig['figure_number']}")
            if para_idx is not None:
                issues.append({'paragraph_index': para_idx, 'comment': f"【插图格式】图{fig['figure_number']}：{issue_text}"})

    for tab in results.get('tables', []):
        for issue_text in tab['info']['issues']:
            para_idx = _find_paragraph_by_text(doc, f"表{tab['table_number']}")
            if para_idx is not None:
                issues.append({'paragraph_index': para_idx, 'comment': f"【表格格式】表{tab['table_number']}：{issue_text}"})

    for formula in results.get('formulas', []):
        for issue_text in formula['info']['issues']:
            if formula.get('formula_number'):
                para_idx = _find_paragraph_by_text(doc, formula['formula_number'])
            else:
                para_idx = None
            if para_idx is not None:
                issues.append({'paragraph_index': para_idx, 'comment': f"【公式格式】{formula.get('formula_number', '')}：{issue_text}"})

    return issues


def collect_header_footer_errors(docx_path):
    """收集页眉页脚检测错误"""
    from .header_footer_checker import HeaderFooterChecker
    checker = HeaderFooterChecker()
    result = checker.check(docx_path)
    issues = []

    for issue_text in result.get('issues', []):
        issues.append({'paragraph_index': 0, 'comment': f"【页眉页脚】{issue_text}"})

    return issues


def collect_reference_errors(docx_path):
    """收集参考文献检测错误"""
    from .reference_checker import GraduateReferenceChecker
    checker = GraduateReferenceChecker()
    results = checker.check_references(docx_path)
    issues = []

    ref_start = checker._find_ref_section(Document(docx_path))
    if ref_start >= 0:
        for it in results.get('title_issues', []):
            if it.startswith('[错误]') or it.startswith('[警告]'):
                issues.append({'paragraph_index': ref_start, 'comment': f"【参考文献标题】{it}"})

    for it in results.get('numbering_issues', []):
        if it.startswith('[警告]'):
            m = re.search(r'\[(\d+)\]', it)
            if m:
                para_idx = _find_ref_paragraph(results['refs'], int(m.group(1)))
                if para_idx is not None:
                    issues.append({'paragraph_index': para_idx, 'comment': f"【参考文献编号】{it}"})

    for category, prefix in [('format_issues', '参考文献格式'), ('punct_issues', '参考文献标点'),
                              ('author_issues', '参考文献作者'), ('type_issues', '参考文献类型')]:
        for it in results.get(category, []):
            m = re.match(r'\[(\d+)\]', it)
            if m:
                para_idx = _find_ref_paragraph(results['refs'], int(m.group(1)))
                if para_idx is not None:
                    desc = re.sub(r'^\[\d+\]\s*', '', it)
                    issues.append({'paragraph_index': para_idx, 'comment': f"【{prefix}】{desc}"})

    return issues


def collect_structure_errors(docx_path):
    """收集论文结构检测错误"""
    from .structure_checker import ThesisStructureChecker
    checker = ThesisStructureChecker()
    doc = Document(docx_path)
    toc_entries = checker.extract_toc_entries(doc)
    order_result = checker.check_order(toc_entries)
    page_result = checker.check_page_continuity(toc_entries)
    issues = []

    for it in order_result.get('issues', []):
        issues.append({'paragraph_index': 0, 'comment': f"【论文结构】{it}"})
    for it in page_result.get('issues', []):
        issues.append({'paragraph_index': 0, 'comment': f"【页码连续性】{it}"})

    return issues


def collect_toc_errors(docx_path):
    """收集目录格式检测错误"""
    from .toc_checker import TocChecker
    checker = TocChecker()
    result = checker.check_toc_format(docx_path)
    issues = []

    if not result['is_autogenerated']:
        issues.append({'paragraph_index': 0, 'comment': "【目录】目录不是自动生成的，建议使用 Word 自动目录功能"})

    for item in result.get('toc_items', []):
        para_idx = item.get('paragraph_index')
        if para_idx is not None:
            for issue_text in result.get('issues', []):
                if item.get('title', '') in issue_text or item.get('number', '') in issue_text:
                    issues.append({'paragraph_index': para_idx, 'comment': f"【目录】{issue_text}"})

    matched_issues = {i['comment'] for i in issues}
    for issue_text in result.get('issues', []):
        full_comment = f"【目录】{issue_text}"
        if full_comment not in matched_issues:
            issues.append({'paragraph_index': 0, 'comment': full_comment})

    return issues


# ============================================================
# 统一入口
# ============================================================

def run_all_checks_with_comments(docx_path, output_path=None):
    """运行所有检测并添加批注。

    Args:
        docx_path: 输入 Word 文档路径
        output_path: 输出路径（默认在原文件名后加 _批注版）

    Returns:
        dict: 各模块检出的错误数量统计
    """
    if not os.path.exists(docx_path):
        print(f"[错误] 文件不存在：{docx_path}")
        return None

    if output_path is None:
        base = os.path.splitext(docx_path)[0]
        output_path = f"{base}_批注版.docx"

    print(f"📄 文档：{docx_path}")
    print(f"📝 输出：{output_path}")
    print()

    reset_comments(author="格式检测器")
    doc = Document(docx_path)

    all_issues = []
    stats = {}

    checkers = [
        ("致谢 & 附录", collect_ack_appendix_errors),
        ("正文格式", collect_body_text_errors),
        ("章节标题", collect_chapter_title_errors),
        ("承诺书", collect_commitment_errors),
        ("封面格式", collect_cover_page_errors),
        ("图表公式", collect_figure_table_errors),
        ("页眉页脚", collect_header_footer_errors),
        ("参考文献", collect_reference_errors),
        ("论文结构", collect_structure_errors),
        ("目录格式", collect_toc_errors),
    ]

    for name, collector in checkers:
        print(f"🔍 检测: {name} ...")
        try:
            issues = collector(docx_path)
            stats[name] = len(issues)
            all_issues.extend(issues)
            if issues:
                print(f"   [警告] 发现 {len(issues)} 个问题")
            else:
                print(f"   [通过] 通过")
        except Exception as e:
            stats[name] = f"错误: {e}"
            print(f"   [错误] 检测出错: {e}")

    print()

    print(f"📝 共收集 {len(all_issues)} 个问题，正在添加批注...")
    count = collect_and_add_comments(doc, all_issues)
    print(f"   成功添加 {count} 条批注")

    print(f"\n💾 保存中...")
    save_with_comments(doc, output_path)
    print(f"[通过] 已保存至: {output_path}")

    print("\n" + "=" * 50)
    print("📊 检测统计")
    print("=" * 50)
    total = 0
    for name, cnt in stats.items():
        if isinstance(cnt, int):
            total += cnt
            status = f"{cnt} 个问题" if cnt > 0 else "通过"
        else:
            status = cnt
        print(f"  {name}: {status}")
    print(f"\n  总计: {total} 个问题 → {count} 条批注")

    return stats


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("用法: python comment_utils.py <docx文件路径> [输出路径]")
        print()
        print("示例:")
        print("  python comment_utils.py 论文.docx")
        print("  python comment_utils.py 论文.docx 输出_批注版.docx")
        sys.exit(1)

    docx_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else None
    run_all_checks_with_comments(docx_path, output_path)

"""
图表公式术语综合检测模块
功能：
1. 检查插图格式（图题、中英文对照）
2. 检查表格格式（表题、中英文对照）
3. 检查公式格式（编号、位置）
4. 检查术语缩略词（定义完整性）
"""

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import re
import os


class FigureTableFormulaTermChecker:
    """图表公式术语综合检测器"""
    
    def __init__(self):
        # 图题标准格式（五号=10.5pt）
        self.figure_caption_format = {
            'chinese_font': ['宋体', 'SimSun'],
            'english_font': ['Times New Roman'],
            'font_size': 10.5,
            'bold': True,
            'alignment': 'CENTER',
            'line_spacing': 1.25
        }
        
        # 表题标准格式（五号=10.5pt）
        self.table_caption_format = {
            'chinese_font': ['宋体', 'SimSun'],
            'english_font': ['Times New Roman'],
            'font_size': 10.5,
            'bold': True,
            'alignment': 'CENTER',
            'line_spacing': 1.25
        }
        
        # 公式编号标准格式
        self.formula_format = {
            'alignment': 'CENTER',
            'number_alignment': 'RIGHT',
            'font_size': 12.0,
            'number_format': r'\((\d+)-(\d+)\)'
        }
        
        # 缩略词模式
        self.acronym_pattern = r'\b[A-Z]{2,10}\b'
        self.definition_pattern = r'([^(（]+)[\(（]([^)）]+)[)）]'
    
    def get_font_name(self, run):
        """获取运行块的字体名称"""
        has_chinese = any('\u4e00' <= char <= '\u9fff' for char in run.text)
        
        try:
            if hasattr(run._element, 'rPr') and run._element.rPr is not None:
                rpr = run._element.rPr
                rfonts_elem = rpr.find(qn('w:rFonts'))
                if rfonts_elem is not None:
                    if has_chinese:
                        eastasia = rfonts_elem.get(qn('w:eastAsia'))
                        if eastasia:
                            return eastasia
                    else:
                        ascii_font = rfonts_elem.get(qn('w:ascii'))
                        if ascii_font:
                            return ascii_font
                    hAnsi = rfonts_elem.get(qn('w:hAnsi'))
                    if hAnsi:
                        return hAnsi
        except:
            pass
        
        if run.font.name:
            return run.font.name
        
        return None
    
    def get_font_size(self, run):
        """获取字号（磅）"""
        if run.font.size:
            return run.font.size.pt
        return None
    
    def get_alignment(self, paragraph):
        """获取对齐方式"""
        alignment_map = {
            0: 'LEFT',
            1: 'CENTER',
            2: 'RIGHT',
            3: 'JUSTIFY',
            4: 'DISTRIBUTE',
            None: 'DEFAULT'
        }
        return alignment_map.get(paragraph.alignment, 'UNKNOWN')
    
    def get_line_spacing(self, paragraph):
        """获取行距（倍数）"""
        try:
            pPr = paragraph._element.find(qn('w:pPr'))
            if pPr is not None:
                spacing = pPr.find(qn('w:spacing'))
                if spacing is not None:
                    line_val = spacing.get(qn('w:line'))
                    if line_val is not None:
                        line_val = int(line_val)
                        # 统一处理：根据值的范围判断编码方式
                        if line_val > 1000:
                            return line_val / 12700.0    # EMU: 15875 → 1.25
                        elif line_val > 100:
                            return line_val / 240.0       # 标准240ths: 300 → 1.25
                        else:
                            return line_val / 20.0        # 非标准×20: 25 → 1.25
        except:
            pass
        return None
    
    # ============ 插图检测 ============
    
    def _is_caption_text(self, text, prefix):
        """判断以 '图X-Y' 或 '表X-Y' 开头的文本是否真的是图题/表题，而非正文引用
        
        区分逻辑：
        - 图题/表题是短标题，如 "图2-4 uniCloud平台架构"
        - 正文引用会带 "中" 字或句子较长，如 "图2-4中展示了..."
        """
        # 去掉开头的 "图X-Y" 或 "表X-Y" 部分
        remainder = re.sub(r'^' + prefix + r'\s*(\d+)[-\.](\d+)\s*', '', text).strip()

        # 如果 "图X-Y" 后紧跟 "中"，说明是正文引用（如 "图2-4中展示了..."）
        if remainder and remainder[0] == '中':
            return False

        # 如果去掉编号后剩余内容很长（超过50字符），大概率是正文段落而非图题
        if len(remainder) > 50:
            return False

        return True

    def identify_figure_caption(self, paragraph):
        """识别图题"""
        text = paragraph.text.strip()
        match = re.match(r'^图\s*(\d+)[-\.](\d+)', text, re.IGNORECASE)
        if match:
            # 排除正文段落中对图的引用
            if not self._is_caption_text(text, '图'):
                return False, None, None, None
            chapter = int(match.group(1))
            seq = int(match.group(2))
            return True, f"{chapter}-{seq}", chapter, seq
        return False, None, None, None
    
    def check_figure_caption_format(self, paragraph, english_caption=None):
        """检查图题格式"""
        text = paragraph.text.strip()
        info = {
            'text': text,
            'alignment': self.get_alignment(paragraph),
            'line_spacing': self.get_line_spacing(paragraph),
            'fonts': [],
            'font_sizes': [],
            'bold_status': [],
            'has_english': False,
            'english_text': english_caption,
            'issues': []
        }
        
        # 检查对齐方式
        if info['alignment'] not in ['CENTER', 'DEFAULT']:
            info['issues'].append(f"对齐方式应为居中，当前为：{info['alignment']}")
        
        # 检查行距
        if info['line_spacing'] and abs(info['line_spacing'] - self.figure_caption_format['line_spacing']) > 0.1:
            info['issues'].append(f"行距应为1.25倍，当前为：{info['line_spacing']}")
        
        # 检查是否有英文对照
        if english_caption:
            info['has_english'] = True
        elif 'Figure' in text or 'figure' in text:
            info['has_english'] = True
        else:
            info['issues'].append("缺少英文对照")
        
        # 检查字体和字号
        for run in paragraph.runs:
            run_text = run.text.strip()
            if not run_text:
                continue
            
            font_size = self.get_font_size(run)
            bold = run.bold
            
            if font_size:
                info['font_sizes'].append(font_size)
            if bold is not None:
                info['bold_status'].append(bold)
            
            # 检查字号
            if font_size and abs(font_size - self.figure_caption_format['font_size']) > 0.5:
                info['issues'].append(f"字号应为五号（10.5pt），当前为：{font_size:.1f}pt")
            
            # 检查加粗
            if bold is False:
                info['issues'].append("应加粗")
            
            break
        
        info['issues'] = list(dict.fromkeys(info['issues']))
        return info
    
    # ============ 表格检测 ============
    
    def identify_table_caption(self, paragraph):
        """识别表题"""
        text = paragraph.text.strip()
        match = re.match(r'^表\s*(\d+)[-\.](\d+)', text, re.IGNORECASE)
        if match:
            # 排除正文段落中对表的引用
            if not self._is_caption_text(text, '表'):
                return False, None, None, None
            chapter = int(match.group(1))
            seq = int(match.group(2))
            return True, f"{chapter}-{seq}", chapter, seq
        return False, None, None, None
    
    def check_table_caption_format(self, paragraph, english_caption=None):
        """检查表题格式"""
        text = paragraph.text.strip()
        info = {
            'text': text,
            'alignment': self.get_alignment(paragraph),
            'line_spacing': self.get_line_spacing(paragraph),
            'fonts': [],
            'font_sizes': [],
            'bold_status': [],
            'has_english': False,
            'english_text': english_caption,
            'issues': []
        }
        
        # 检查对齐方式
        if info['alignment'] not in ['CENTER', 'DEFAULT']:
            info['issues'].append(f"对齐方式应为居中，当前为：{info['alignment']}")
        
        # 检查行距
        if info['line_spacing'] and abs(info['line_spacing'] - self.table_caption_format['line_spacing']) > 0.1:
            info['issues'].append(f"行距应为1.25倍，当前为：{info['line_spacing']}")
        
        # 检查是否有英文对照
        if english_caption:
            info['has_english'] = True
        elif 'Table' in text or 'table' in text:
            info['has_english'] = True
        else:
            info['issues'].append("缺少英文对照")
        
        # 检查字体和字号
        for run in paragraph.runs:
            run_text = run.text.strip()
            if not run_text:
                continue
            
            font_size = self.get_font_size(run)
            bold = run.bold
            
            if font_size:
                info['font_sizes'].append(font_size)
            if bold is not None:
                info['bold_status'].append(bold)
            
            if font_size and abs(font_size - self.table_caption_format['font_size']) > 0.5:
                info['issues'].append(f"字号应为五号（10.5pt），当前为：{font_size:.1f}pt")
            
            if bold is False:
                info['issues'].append("应加粗")
            
            break
        
        info['issues'] = list(dict.fromkeys(info['issues']))
        return info
    
    # ============ 公式检测 ============
    
    def identify_formula_number(self, paragraph):
        """识别公式编号"""
        text = paragraph.text.strip()
        match = re.search(r'\((\d+)[-\.](\d+)\)', text)
        if match:
            chapter = int(match.group(1))
            seq = int(match.group(2))
            number = f"({chapter}-{seq})"
            return True, number, chapter, seq
        return False, None, None, None
    
    def check_formula_paragraph(self, paragraph):
        """检查包含公式编号的段落"""
        text = paragraph.text.strip()
        info = {
            'text': text,
            'alignment': self.get_alignment(paragraph),
            'has_formula': False,
            'formula_number': None,
            'issues': []
        }
        
        # 检查是否包含公式对象
        try:
            omath_elements = paragraph._element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
            if omath_elements:
                info['has_formula'] = True
        except:
            pass
        
        # 检查是否包含公式编号
        has_number, number, chapter, seq = self.identify_formula_number(paragraph)
        if has_number:
            info['formula_number'] = number
            
            # 检查编号格式（半角括号）
            if re.search(r'（\d+[-\.]\d+）', text):
                info['issues'].append("公式编号应使用半角括号()，当前使用了全角括号（）")
        
        return info
    
    # ============ 术语缩略词检测 ============
    
    def extract_acronyms(self, text):
        """提取缩略词"""
        return re.findall(self.acronym_pattern, text)
    
    def extract_definitions(self, text):
        """提取缩略词定义"""
        matches = re.findall(self.definition_pattern, text)
        definitions = []
        
        for chinese_part, english_part in matches:
            chinese_part = chinese_part.strip()
            english_part = english_part.strip()
            
            acronym_match = re.search(r'\b([A-Z]{2,10})\b', english_part)
            if acronym_match:
                acronym = acronym_match.group(1)
                definitions.append({
                    'chinese': chinese_part,
                    'english': english_part,
                    'acronym': acronym
                })
        
        return definitions
    
    # ============ 综合检测 ============
    
    def check_all(self, docx_path):
        """检查文档中的所有图表公式术语"""
        doc = Document(docx_path)
        
        results = {
            'figures': [],
            'tables': [],
            'formulas': [],
            'acronyms': {
                'defined': {},
                'undefined': {},
                'common': {},
                'definition_examples': []
            }
        }
        
        in_body = False
        formula_index = 0
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            
            if not text:
                continue
            
            # 进入正文区域
            if not in_body and re.match(r'^第[1一]章', text):
                if not re.search(r'\s+\d+$', text):
                    in_body = True
            
            if not in_body:
                continue
            
            # 检测正文结束
            if re.match(r'^参\s*考\s*文\s*献', text) or re.match(r'^致\s*谢', text):
                break
            
            # 检查图题
            is_figure, fig_num, chapter, seq = self.identify_figure_caption(para)
            if is_figure:
                english_caption = None
                if i + 1 < len(doc.paragraphs):
                    next_text = doc.paragraphs[i + 1].text.strip()
                    if re.match(rf'Figure\s*{chapter}[-\.]{seq}', next_text, re.IGNORECASE):
                        english_caption = next_text
                
                caption_info = self.check_figure_caption_format(para, english_caption)
                results['figures'].append({
                    'figure_number': fig_num,
                    'chapter': chapter,
                    'sequence': seq,
                    'info': caption_info,
                    'english_caption': english_caption,
                    'paragraph_index': i
                })
            
            # 检查表题
            is_table, tab_num, chapter, seq = self.identify_table_caption(para)
            if is_table:
                english_caption = None
                if i + 1 < len(doc.paragraphs):
                    next_text = doc.paragraphs[i + 1].text.strip()
                    if re.match(rf'Table\s*{chapter}[-\.]{seq}', next_text, re.IGNORECASE):
                        english_caption = next_text
                
                caption_info = self.check_table_caption_format(para, english_caption)
                results['tables'].append({
                    'table_number': tab_num,
                    'chapter': chapter,
                    'sequence': seq,
                    'info': caption_info,
                    'english_caption': english_caption,
                    'paragraph_index': i
                })
            
            # 检查公式
            formula_info = self.check_formula_paragraph(para)
            if formula_info['has_formula'] or formula_info['formula_number']:
                formula_index += 1
                has_number, number, chapter, seq = self.identify_formula_number(para)
                results['formulas'].append({
                    'index': formula_index,
                    'formula_number': number if has_number else None,
                    'chapter': chapter if has_number else None,
                    'sequence': seq if has_number else None,
                    'info': formula_info,
                    'paragraph_index': i
                })
            
            # 检查术语缩略词
            definitions = self.extract_definitions(text)
            for defn in definitions:
                acronym = defn['acronym']
                if acronym not in results['acronyms']['defined']:
                    results['acronyms']['defined'][acronym] = {
                        'definition': defn,
                        'text_preview': text[:100] + '...' if len(text) > 100 else text
                    }
                    results['acronyms']['definition_examples'].append({
                        'acronym': acronym,
                        'chinese': defn['chinese'],
                        'english': defn['english'],
                        'text_preview': text[:100] + '...' if len(text) > 100 else text
                    })
            
            acronyms = self.extract_acronyms(text)
            for acronym in acronyms:
                if acronym in [
                    'A', 'I', 'OK', 'ID', 'VS', 'IT',
                    # 常见技术缩略词，不需要定义
                    'XML', 'HTML', 'CSS', 'DOCX', 'DOC', 'DOT', 'PPT', 'PPTX', 'XLS', 'XLSX',
                    'PDF', 'URL', 'URI', 'HTTP', 'HTTPS', 'FTP', 'API', 'SQL', 'JSON',
                    'CPU', 'GPU', 'RAM', 'ROM', 'SSD', 'HDD',
                    'GUI', 'CLI', 'IDE', 'SDK', 'JDK',
                    'TCP', 'UDP', 'IP', 'DNS', 'DHCP', 'LAN', 'WAN',
                    'UTF', 'ASCII', 'Unicode',
                    'OS', 'BIOS', 'UEFI', 'USB', 'WiFi', 'IoT',
                    'AI', 'ML', 'DL', 'NLP', 'CV',
                    'ERP', 'CRM', 'OA', 'MIS',
                    'HTML5', 'CSS3', 'ES6',
                ]:
                    # 记录常见缩略词，不算问题
                    if 'common' not in results['acronyms']:
                        results['acronyms']['common'] = {}
                    if acronym not in results['acronyms']['common']:
                        results['acronyms']['common'][acronym] = []
                    results['acronyms']['common'][acronym].append({
                        'text_preview': text[:100] + '...' if len(text) > 100 else text
                    })
                    continue
                if acronym in results['acronyms']['defined']:
                    continue
                if acronym not in results['acronyms']['undefined']:
                    results['acronyms']['undefined'][acronym] = []
                results['acronyms']['undefined'][acronym].append({
                    'text_preview': text[:100] + '...' if len(text) > 100 else text
                })
        
        # 排序
        results['figures'].sort(key=lambda x: (x['chapter'], x['sequence']))
        results['tables'].sort(key=lambda x: (x['chapter'], x['sequence']))
        if results['formulas']:
            numbered = [f for f in results['formulas'] if f['formula_number']]
            if numbered:
                numbered.sort(key=lambda x: (x['chapter'], x['sequence']))
        
        return results
    
    def highlight_paragraph(self, paragraph):
        """在段落中添加高亮背景"""
        for run in paragraph.runs:
            run.font.highlight_color = 7  # 7 表示黄色

    def generate_report(self, docx_path, report_path=None, highlight_path=None):
        """生成综合检测报告"""
        print("开始图表公式术语综合检测...")

        results = self.check_all(docx_path)

        # 如果提供了高亮路径，对有问题的段落进行高亮
        if highlight_path:
            try:
                hl_doc = Document(highlight_path)

                # 高亮有问题的图题段落
                for fig in results.get('figures', []):
                    if fig['info']['issues']:
                        pi = fig['paragraph_index']
                        if 0 <= pi < len(hl_doc.paragraphs):
                            self.highlight_paragraph(hl_doc.paragraphs[pi])

                # 高亮有问题的表题段落
                for tab in results.get('tables', []):
                    if tab['info']['issues']:
                        pi = tab['paragraph_index']
                        if 0 <= pi < len(hl_doc.paragraphs):
                            self.highlight_paragraph(hl_doc.paragraphs[pi])

                # 高亮有问题的公式段落
                for formula in results.get('formulas', []):
                    if formula['info']['issues']:
                        pi = formula['paragraph_index']
                        if 0 <= pi < len(hl_doc.paragraphs):
                            self.highlight_paragraph(hl_doc.paragraphs[pi])

                hl_doc.save(highlight_path)
                print(f"图表公式问题已高亮至: {highlight_path}")
            except Exception as e:
                print(f"高亮图表公式问题失败: {e}")
        
        # 生成Markdown报告
        report_lines = []
        report_lines.append("# 图表公式术语综合检测报告\n")
        
        # 基本信息
        report_lines.append("## 基本信息")
        report_lines.append(f"- 文档路径: {docx_path}")
        report_lines.append(f"- 插图数量: {len(results['figures'])}个")
        report_lines.append(f"- 表格数量: {len(results['tables'])}个")
        report_lines.append(f"- 公式数量: {len(results['formulas'])}个")
        report_lines.append(f"- 缩略词数量: {len(results['acronyms']['defined']) + len(results['acronyms'].get('undefined', {}))}个")
        report_lines.append(f"  - 有定义: {len(results['acronyms']['defined'])}个")
        report_lines.append(f"  - 可能缺少定义: {len(results['acronyms'].get('undefined', {}))}个")
        common_count = len(results['acronyms'].get('common', {}))
        if common_count > 0:
            report_lines.append(f"  - 常见缩略词（免检）: {common_count}个")
        report_lines.append("")
        
        # 统计问题
        figure_issues = sum(len(f['info']['issues']) for f in results['figures'])
        table_issues = sum(len(t['info']['issues']) for t in results['tables'])
        formula_issues = sum(len(f['info']['issues']) for f in results['formulas'])
        
        report_lines.append("## 检测概况\n")
        report_lines.append(f"### 插图检测")
        if figure_issues == 0:
            report_lines.append(f"[通过] {len(results['figures'])}个插图全部格式规范\n")
        else:
            report_lines.append(f"[警告] 发现 {figure_issues} 个问题\n")
        
        report_lines.append(f"### 表格检测")
        if table_issues == 0:
            report_lines.append(f"[通过] {len(results['tables'])}个表格全部格式规范\n")
        else:
            report_lines.append(f"[警告] 发现 {table_issues} 个问题\n")
        
        report_lines.append(f"### 公式检测")
        if formula_issues == 0:
            report_lines.append(f"[通过] {len(results['formulas'])}个公式全部格式规范\n")
        else:
            report_lines.append(f"[警告] 发现 {formula_issues} 个问题\n")
        
        report_lines.append(f"### 术语缩略词")
        undefined_acronyms = results['acronyms'].get('undefined', {})
        undefined_count = len(undefined_acronyms)
        common_acronyms = results['acronyms'].get('common', {})
        common_count = len(common_acronyms)
        if undefined_count == 0:
            report_lines.append(f"[通过] 所有缩略词都有定义")
            if common_count > 0:
                report_lines.append(f"（另有 {common_count} 个常见缩略词免检）")
            report_lines.append("")
        else:
            report_lines.append(f"[信息] 发现 {undefined_count} 个未定义缩略词（仅供参考，不计入问题）\n")
        
        # 详细结果
        if figure_issues > 0:
            report_lines.append("## 插图问题详情\n")
            for fig in results['figures']:
                if fig['info']['issues']:
                    report_lines.append(f"### 图{fig['figure_number']}\n")
                    report_lines.append(f"- **内容**: {fig['info']['text']}")
                    report_lines.append(f"- **问题**:")
                    for issue in fig['info']['issues']:
                        report_lines.append(f"  - {issue}")
                    report_lines.append("")
        
        if table_issues > 0:
            report_lines.append("## 表格问题详情\n")
            for tab in results['tables']:
                if tab['info']['issues']:
                    report_lines.append(f"### 表{tab['table_number']}\n")
                    report_lines.append(f"- **内容**: {tab['info']['text']}")
                    report_lines.append(f"- **问题**:")
                    for issue in tab['info']['issues']:
                        report_lines.append(f"  - {issue}")
                    report_lines.append("")
        
        if formula_issues > 0:
            report_lines.append("## 公式问题详情\n")
            for formula in results['formulas']:
                if formula['info']['issues']:
                    report_lines.append(f"### 公式 {formula['formula_number'] or formula['index']}\n")
                    report_lines.append(f"- **问题**:")
                    for issue in formula['info']['issues']:
                        report_lines.append(f"  - {issue}")
                    report_lines.append("")
        
        if undefined_count > 0:
            report_lines.append("## 未定义缩略词（仅供参考）\n")
            sorted_undefined = sorted(
                undefined_acronyms.items(),
                key=lambda x: len(x[1]),
                reverse=True
            )
            for acronym, occurrences in sorted_undefined[:10]:
                report_lines.append(f"- **{acronym}** ({len(occurrences)}次)")
            if len(sorted_undefined) > 10:
                report_lines.append(f"\n（还有 {len(sorted_undefined) - 10} 个）")
            report_lines.append("")
        
        # 总结（缩略词不计入问题）
        total_issues = figure_issues + table_issues + formula_issues
        report_lines.append("## 总结\n")
        if total_issues == 0:
            report_lines.append("[通过] **所有检测项目全部符合要求！**")
        else:
            report_lines.append(f"**检测发现**:")
            if total_issues > 0:
                report_lines.append(f"- 图表公式格式问题: {total_issues}个")
        
        # 写入报告
        with open(report_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(report_lines))
        
        print(f"\n报告已保存至: {report_path}")
        print(f"\n检测统计:")
        print(f"  插图: {len(results['figures'])}个 ({figure_issues}个问题)")
        print(f"  表格: {len(results['tables'])}个 ({table_issues}个问题)")
        print(f"  公式: {len(results['formulas'])}个 ({formula_issues}个问题)")
        print(f"  缩略词: 已定义{len(results['acronyms']['defined'])}个, 未定义{len(results['acronyms'].get('undefined', {}))}个（仅供参考）")


def main():
    """主函数"""
    import os
    import sys
    current_dir = os.path.dirname(os.path.abspath(__file__))
    base_dir = os.path.dirname(current_dir)
    docx_path = os.path.join(base_dir, "测试论文.docx")
    reports_dir = os.path.join(base_dir, "reports")
    os.makedirs(reports_dir, exist_ok=True)
    report_path = os.path.join(reports_dir, "综合检测报告.md")
    
    if not os.path.exists(docx_path):
        print(f"[错误] 文件不存在：{docx_path}")
        print(f"   当前目录：{current_dir}")
        sys.exit(1)
    
    try:
        checker = FigureTableFormulaTermChecker()
        checker.generate_report(docx_path, report_path)
    except Exception as e:
        print(f"[错误] 检测过程中出错：{str(e)}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()

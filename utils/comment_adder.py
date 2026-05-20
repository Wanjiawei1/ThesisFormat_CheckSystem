"""
批注添加模块
功能：
1. 复制原始Word文档
2. 在检测到的错误位置添加批注
3. 通过XML注入实现批注功能
"""

import shutil
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
from datetime import datetime


class CommentAdder:
    """批注添加器"""
    
    def __init__(self):
        pass
    
    def copy_file(self, original_file, target_file):
        """复制文件"""
        try:
            shutil.copyfile(original_file, target_file)
            return True
        except Exception as e:
            print(f"复制文件失败: {str(e)}")
            return False
    
    def add_comment_to_paragraph(self, paragraph, comment_text, author="格式检测助手"):
        """在段落中添加批注"""
        # 获取文档对象
        doc = paragraph._parent._parent
        
        # 获取或创建comments部件
        if not hasattr(doc, '_comments_part'):
            # 创建comments部件
            comments_part = doc.part.add_new_part('application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml')
            comments_part._element = OxmlElement('w:comments')
            comments_part._element.set(qn('xmlns:w'), 'http://schemas.openxmlformats.org/wordprocessingml/2006/main')
            doc._comments_part = comments_part
        else:
            comments_part = doc._comments_part
        
        # 生成唯一ID
        comment_id = str(len(comments_part._element) + 1)
        
        # 创建comment元素
        comment = OxmlElement('w:comment')
        comment.set(qn('w:id'), comment_id)
        comment.set(qn('w:author'), author)
        comment.set(qn('w:date'), datetime.now().isoformat())
        
        # 添加批注内容
        p = OxmlElement('w:p')
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = comment_text
        r.append(t)
        p.append(r)
        comment.append(p)
        
        # 将comment添加到comments部件
        comments_part._element.append(comment)
        
        # 在段落前添加commentRangeStart
        comment_range_start = OxmlElement('w:commentRangeStart')
        comment_range_start.set(qn('w:id'), comment_id)
        paragraph._p.insert(0, comment_range_start)
        
        # 在段落末尾添加commentRangeEnd
        comment_range_end = OxmlElement('w:commentRangeEnd')
        comment_range_end.set(qn('w:id'), comment_id)
        paragraph._p.append(comment_range_end)
        
        # 添加commentReference
        comment_reference = OxmlElement('w:commentReference')
        comment_reference.set(qn('w:id'), comment_id)
        r = OxmlElement('w:r')
        r.append(comment_reference)
        paragraph._p.append(r)
    
    def add_comments_to_issues(self, original_file, issues, output_file):
        """根据检测结果添加批注
        
        Args:
            original_file: 原始文件路径
            issues: 检测到的问题列表，每个问题包含paragraph_index和issues字段
            output_file: 输出文件路径
        """
        # 复制文件
        if not self.copy_file(original_file, output_file):
            return False
        
        try:
            # 打开副本
            doc = Document(output_file)
            
            # 遍历所有问题
            for issue in issues:
                paragraph_index = issue.get('paragraph_index', 0)
                issue_list = issue.get('issues', [])
                
                # 确保段落索引有效
                if 0 < paragraph_index <= len(doc.paragraphs):
                    # 获取目标段落
                    target_paragraph = doc.paragraphs[paragraph_index - 1]
                    
                    # 构建批注文本
                    comment_text = "格式问题:\n"
                    for i, problem in enumerate(issue_list, 1):
                        comment_text += f"{i}. {problem}\n"
                    
                    # 添加批注
                    self.add_comment_to_paragraph(target_paragraph, comment_text)
            
            # 保存修改
            doc.save(output_file)
            return True
        except Exception as e:
            print(f"添加批注失败: {str(e)}")
            return False
